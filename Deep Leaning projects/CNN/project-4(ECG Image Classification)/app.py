from flask import Flask, render_template, request, redirect, url_for, session, flash, send_from_directory
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from tensorflow.keras.models import load_model
from tensorflow.keras.preprocessing import image
from functools import wraps
from datetime import datetime
import numpy as np
import os

app = Flask(__name__)
app.secret_key = "change-this-in-production"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, 'users.xlsx')
PENDING_FILE = os.path.join(BASE_DIR, 'pending_users.xlsx')
LOG_FILE = os.path.join(BASE_DIR, 'logins.xlsx')
MODEL_PATH = os.path.join(BASE_DIR, 'ecg_cnn_cpu_model.h5')
STATIC_DIR = os.path.join(BASE_DIR, 'static')
UPLOAD_FOLDER = os.path.join(STATIC_DIR, 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

if not os.path.exists(MODEL_PATH):
    print("[WARN] Model file 'ecg_cnn_cpu_model.h5' not found. Place it at project root for predictions.")
    model = None
else:
    model = load_model(MODEL_PATH)

class_names = ['Myocardial_Infarction', 'History_MI', 'Abnormal_Heartbeat', 'Normal']

def _init_users_sheet():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Users"
        ws.append(["Username", "PasswordHash", "Email", "Role"])
        wb.save(EXCEL_FILE)

    wb = load_workbook(EXCEL_FILE)
    ws = wb["Users"]

    have_admin = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 4 and (row[3] or "").strip().lower() == "admin":
            have_admin = True
            break
    if not have_admin:
        ws.append(["admin", generate_password_hash("admin123"), "admin@example.com", "Admin"])
        wb.save(EXCEL_FILE)
        print("[INFO] Seeded default admin: username=admin password=admin123")

def _init_pending_sheet():
    if not os.path.exists(PENDING_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Pending"
        ws.append(["Username", "Password", "Email", "RequestedAt"])
        wb.save(PENDING_FILE)

def _init_logins_sheet():
    if not os.path.exists(LOG_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Logins"
        ws.append(["Username", "Role", "LoginTime"])
        wb.save(LOG_FILE)

_init_users_sheet()
_init_pending_sheet()
_init_logins_sheet()

def _read_users():
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Users"]
    users = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        username, pwd_hash, email, role = (row + (None, None, None, None))[:4]
        if username and pwd_hash:
            users[str(username).strip()] = {
                "password": str(pwd_hash).strip(),
                "email": str(email or "").strip(),
                "role": str(role or "User").strip()
            }
    return users

def _add_user(username: str, password: str, email: str, role: str = "User"):
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Users"]
    ws.append([username, generate_password_hash(password), email, role])
    wb.save(EXCEL_FILE)

def _add_pending_user(username: str, password: str, email: str):
    wb = load_workbook(PENDING_FILE)
    ws = wb["Pending"]
    ws.append([username, password, email, datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    wb.save(PENDING_FILE)

def _read_pending():
    wb = load_workbook(PENDING_FILE)
    ws = wb["Pending"]
    pending = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        username, password, email, requested = (row + (None, None, None, None))[:4]
        if username:
            pending.append({"username": username, "password": password, "email": email, "requested": requested})
    return pending

def _delete_pending(username: str):
    wb = load_workbook(PENDING_FILE)
    ws = wb["Pending"]
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if (row[0].value or "").strip() == username:
            ws.delete_rows(i)
            break
    wb.save(PENDING_FILE)

def _log_login(username: str, role: str):
    wb = load_workbook(LOG_FILE)
    ws = wb["Logins"]
    ws.append([username, role, datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    wb.save(LOG_FILE)

def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if "username" not in session:
            flash("Please login to continue.", "error")
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapper

def admin_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if "username" not in session:
            flash("Please login to continue.", "error")
            return redirect(url_for("login"))
        if session.get("role") != "Admin":
            flash("Admins only.", "error")
            return redirect(url_for("user_dashboard"))
        return view(*args, **kwargs)
    return wrapper

@app.route('/')
def root():
    if "username" in session:
        return redirect(url_for('admin_dashboard' if session.get("role") == "Admin" else 'user_dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')

        users = _read_users()
        if username in users and check_password_hash(users[username]["password"], password):
            session['username'] = username
            session['role'] = users[username]["role"]
            _log_login(username, session["role"])
            return redirect(url_for('admin_dashboard' if session["role"] == "Admin" else 'user_dashboard'))
        else:
            for p in _read_pending():
                if p["username"] == username:
                    flash("Your registration is pending admin approval.", "error")
                    break
            else:
                flash('Invalid username or password.', 'error')
            return redirect(url_for('login'))

    return render_template('login.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email    = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        confirm  = request.form.get('confirm', '')

        if not username or not password or not email:
            flash('All fields are required.', 'error')
            return redirect(url_for('register'))
        if password != confirm:
            flash('Passwords do not match.', 'error')
            return redirect(url_for('register'))

        users = _read_users()
        if username in users:
            flash('Username already exists.', 'error')
            return redirect(url_for('register'))

        _add_pending_user(username, password, email)
        flash('Registration request sent. Wait for admin approval.', 'success')
        return redirect(url_for('login'))

    return render_template('register.html')

@app.route('/logout')
@login_required
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/user_dashboard')
@login_required
def user_dashboard():
    if session.get("role") == "Admin":
        return redirect(url_for("admin_dashboard"))
    return render_template('user_dashboard.html', username=session['username'])

@app.route('/admin_dashboard')
@login_required
def admin_dashboard():
    if session.get("role") != "Admin":
        return redirect(url_for("user_dashboard"))
    return render_template('admin_dashboard.html', username=session['username'])

@app.route('/home')
@login_required
def home_page():
    return render_template('home.html', username=session['username'], role=session.get("role"))

@app.route('/classification')
@login_required
def classification_page():
    return render_template('classification.html', username=session['username'], role=session.get("role"))

@app.route('/description')
@login_required
def description_page():
    return render_template('description.html', username=session['username'], role=session.get("role"))

@app.route('/predict', methods=['POST'])
@login_required
def predict():
    if model is None:
        flash("Model not loaded. Place 'ecg_cnn_cpu_model.h5' at project root.", "error")
        return redirect(url_for("classification_page"))

    if 'image' not in request.files or request.files['image'].filename == '':
        flash("No image uploaded.", "error")
        return redirect(url_for("classification_page"))

    file = request.files['image']
    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    img = image.load_img(filepath, target_size=(128, 128))
    arr = image.img_to_array(img) / 255.0
    arr = arr.reshape((1, 128, 128, 3))

    pred = model.predict(arr)
    predicted_class = class_names[int(np.argmax(pred))]

    return render_template(
        "classification.html",
        username=session['username'],
        role=session.get("role"),
        prediction=predicted_class,
        uploaded_image=url_for('uploaded_file', filename=filename)
    )

@app.route('/uploads/<path:filename>')
@login_required
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/pending_users')
@admin_required
def pending_users():
    return render_template('pending_users.html', users=_read_pending())

@app.route('/approve_user/<username>', methods=['POST'])
@admin_required
def approve_user(username):
    pending = _read_pending()
    found = None
    for p in pending:
        if p['username'] == username:
            found = p
            break
    if not found:
        flash("Pending user not found.", "error")
        return redirect(url_for('pending_users'))
    _add_user(found['username'], found['password'], found['email'], role="User")
    _delete_pending(found['username'])
    flash(f"Approved {username}.", "success")
    return redirect(url_for('pending_users'))

@app.route('/reject_user/<username>', methods=['POST'])
@admin_required
def reject_user(username):
    _delete_pending(username)
    flash(f"Rejected {username}.", "success")
    return redirect(url_for('pending_users'))

@app.route('/monitor')
@admin_required
def monitor():
    wb = load_workbook(LOG_FILE)
    ws = wb["Logins"]
    logs = [row for row in ws.iter_rows(min_row=2, values_only=True)]
    return render_template('monitor.html', logs=logs)

if __name__ == "__main__":
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    app.run(debug=True)
